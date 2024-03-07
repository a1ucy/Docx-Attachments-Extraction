#实现一个文档附件提取器，能够对常见的word文档，所有附件数据进行提取，附件包括但不限于文件、压缩包等
import os
import docx
import olefile
from oletools import oleobj
from openpyxl import load_workbook

# doc_path  = input('请输入word文件地址：')
doc_path = 'test.docx'
prev_path = os.path.dirname(doc_path)
output_path = os.path.join(prev_path, 'attachments')

doc = docx.Document(doc_path)
items = doc.part.rels.items()

lis = []
for id, part in items:
    if str(part.target_ref).startswith('embedding'):
        if not os.path.exists(output_path):
            os.makedirs(output_path)
        part_name = str(part.target_ref).replace('embedding','').replace('/','')
        part_path = os.path.join(output_path, part_name)
        lis.append(part_path)
        with open(part_path, 'wb') as f:
            f.write(part.target_part.blob)
            
# extract non-docx or xlsx file attachments
remain = []
for i in lis:
    if oleobj.main([i]) == 1:
        os.remove(i)
    else:
        remain.append(i)

# extract xlsx file
if remain:
    num = 1
    for f in remain:
        if olefile.isOleFile(f):
            content = olefile.OleFileIO(f).openstream("package")
            try:
                book = load_workbook(content)
                excel_name = 'excel'+str(num)+'.xlsx'
                excel_path = os.path.join(output_path, excel_name)
                book.save(excel_path)
                os.remove(f)
                num += 1
                print("Excel saved successfully", num)
            except OSError:
                print('File', os.path.basename(f) , 'unable to extract.')
                continue